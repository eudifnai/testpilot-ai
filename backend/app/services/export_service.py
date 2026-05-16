from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.schemas.testcase import Testcase


class ExportService:
    headers = [
        "Case ID",
        "Module",
        "Title",
        "Precondition",
        "Steps",
        "Test Data",
        "Expected Result",
        "Priority",
        "Case Type",
        "Remark",
    ]

    def export_testcases(self, testcases: list[Testcase]) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Testcases"

        for column_index, header in enumerate(self.headers, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="1D4ED8")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row_index, testcase in enumerate(testcases, start=2):
            worksheet.cell(row=row_index, column=1, value=testcase.case_id)
            worksheet.cell(row=row_index, column=2, value=testcase.module)
            worksheet.cell(row=row_index, column=3, value=testcase.title)
            worksheet.cell(row=row_index, column=4, value=testcase.precondition)
            worksheet.cell(row=row_index, column=5, value="\n".join(testcase.steps))
            worksheet.cell(row=row_index, column=6, value=testcase.test_data)
            worksheet.cell(row=row_index, column=7, value=testcase.expected_result)
            worksheet.cell(row=row_index, column=8, value=testcase.priority)
            worksheet.cell(row=row_index, column=9, value=testcase.case_type)
            worksheet.cell(row=row_index, column=10, value=testcase.remark)

        widths = [12, 18, 28, 26, 40, 24, 34, 12, 16, 20]
        for column_index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[chr(64 + column_index)].width = width

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
